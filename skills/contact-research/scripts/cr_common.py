"""Shared helpers for the contact_research skill: state I/O, validation, redaction, naming."""
import json, os, re, datetime, csv

STAGES = ["init", "hubspot", "associations", "activity", "derived", "research", "rendered"]

CHARACTERIZATION_BLOCKLIST = [
    "personality", "mindset", "demeanor", "prefers", "preference", "responds best",
    "approach him", "approach her", "discreet", "temperament", "engage him by", "engage her by",
    "topics to avoid",
]

RESTRICTED_PATTERNS = [
    (re.compile(r"\b\d{3}-\d{2}-\d{4}\b"), "ssn"),
    (re.compile(r"\b(?:\d[ -]?){13,19}\b"), "card_or_account_number"),
    (re.compile(r"\b(?:acct|account|routing|aba)\b[^\n]{0,30}\b\d{6,}\b", re.I), "account_number"),
    (re.compile(r"\bPIN\b\s*(?:is|was|number|#|[:=])?\s*\d{4,8}\b"), "pin"),
    (re.compile(r"\b(?:wire|ach)\s+(?:instructions|transfer\s+details|routing)\b", re.I), "wire_or_ach_details"),
    (re.compile(r"\b(?:password|passcode|passwd|pwd|pass\s*word)\b\s*(?:is|was|[:=])\s*\S+", re.I), "credential_password"),
    (re.compile(r"\b(?:username|user\s*name|user\s*id|login|log-in)\b\s*[:=]\s*\S+", re.I), "credential_login"),
]
# URLs are removed before the pattern screen runs: press-release IDs (Business Wire), SEC EDGAR accession numbers,
# BBB/LoopNet/LinkedIn/Facebook object IDs are 13–19 digit runs that are not card or account numbers. Five bulk runs
# (about 1,300 contacts) rejected roughly 20 legitimate source URLs per run before this exemption.
URL_RE = re.compile(r"https?://[^\s\"'<>)\]]+", re.I)
# Entries are stems, not whole words: screen_text matches a word boundary followed by the entry, so
# "surger" catches surgery/surgeries and "diagnos" catches diagnosis/diagnosed. Before v1.0 several
# entries were whole words and missed ordinary inflections — "surgery" did not match "surgeries" in a
# note describing a child's operations.
#
# Three terms are deliberately NOT here, each measured against 9,295 real engagement bodies:
#   "stroke"   – 28 matches, every one the CSS property `-webkit-text-stroke-width` in HubSpot HTML.
#   "widow"    – 60 matches, every one the CSS property `widows:`. Use "widowed"/"widower" instead.
#   "cemetery" – matched legitimate grant recipients (Friends of the Cedar Key Cemetery). "buried"
#                carries the bereavement signal without the grant-recipient false positives.
# "memorial service" is a phrase for the same reason: bare "memorial" matches Memorial Day and
# memorial grants. Re-measure before adding any term that could appear in markup or an org name.
HEALTH_HARDSHIP_LEXICON = [
    "cancer", "diagnos", "hospice", "surger", "chemo", "dementia", "alzheim", "depression", "rehab",
    "addict", "bankrupt", "foreclos", "divorc", "terminal", "illness", "disabilit", "medical", "medication",
    # added v1.0 after bulk runs surfaced these in engagement bodies that the earlier list let through
    "bereave", "prayer request", "pray for", "prayers for", "griev", "funeral", "passed away", "pregnan",
    "caregiv", "hospitaliz",
    # added v1.0 after the noneassigned_37 and others_543 runs: injury and bereavement wording the
    # orchestrator had to redact by hand on ten records (fractured wrist, bad ankle, knee recovery,
    # brittle bone disease, a memorial service, a burial place, an illness).
    "memorial service", "fractur", "injur", "broken bone", "brittle bone", "disease", "hardship",
    "ankle", "knee", "wrist", "buried", "sick", "widowed", "widower",
]
# Board Confidential screen (IT15): applied to engagement bodies only when the contact is Signatry staff or a board
# member (see signatry_insider). A donor's note that mentions their own company's board is not Board Confidential.
BOARD_CONFIDENTIAL_LEXICON = [
    "board meeting", "board minutes", "board packet", "board update", "board agenda", "board resolution",
    "executive session", "board only", "privileged", "the board", "board retreat", "board call",
]

# Email domains that are personal mailboxes, privacy relays, or consumer ISPs – never an employer.
# Extended in v1.0 after derive.py classified proton.me, pm.me, duck.com, frontier.com, earthlink.net, and
# similar as corporate across five bulk runs.
PERSONAL_DOMAINS = {
    "gmail.com", "googlemail.com", "yahoo.com", "ymail.com", "yahoo.co.uk", "hotmail.com", "outlook.com", "live.com",
    "msn.com", "aol.com", "icloud.com", "me.com", "mac.com", "protonmail.com", "proton.me", "pm.me", "protonmail.ch",
    "duck.com", "hey.com", "fastmail.com", "fastmail.fm", "tutanota.com", "tuta.io", "hushmail.com", "zoho.com",
    "mail.com", "gmx.com", "gmx.net", "comcast.net", "xfinity.com", "att.net", "sbcglobal.net", "bellsouth.net",
    "pacbell.net", "verizon.net", "frontier.com", "frontiernet.net", "earthlink.net", "mindspring.com", "cox.net",
    "charter.net", "spectrum.net", "rr.com", "roadrunner.com", "twc.com", "optonline.net", "optimum.net",
    "centurylink.net", "centurytel.net", "q.com", "embarqmail.com", "windstream.net", "juno.com", "netzero.net",
    "wavecable.com", "cableone.net", "sky.com", "btinternet.com", "shaw.ca", "rogers.com", "telus.net",
    # added v1.0 after the noneassigned_37 and others_543 runs: regional ISPs, a rural telephone
    # co-op, a consumer mail provider, and a fax-to-email service, each classed as an employer and
    # each producing a company record for a domain with no company behind it.
    "ptd.net", "iamotelephone.com", "rocketmail.com", "ibyfax.com", "iland.net", "bex.net", "reagan.com",
}
# Broker-dealer and wirehouse domains shared by thousands of unrelated advisory practices. A company record keyed
# to one of these describes one practice, not the next contact's employer, so company_prior lookups skip them and
# the research playbook asks for a per-practice slug (the practice's own domain) instead.
SHARED_DOMAINS = {
    "nm.com", "lpl.com", "ml.com", "ms.com", "morganstanley.com", "ubs.com", "wellsfargo.com", "wfadvisors.com",
    "edwardjones.com", "raymondjames.com", "rjf.com", "ampf.com", "ameriprise.com", "stifel.com", "rwbaird.com",
    "schwab.com", "fidelity.com", "fmr.com", "jpmorgan.com", "jpmchase.com", "bofa.com", "truist.com", "suntrust.com",
    "cetera.com", "ceterawealth.com", "osaic.com", "commonwealth.com", "securian.com", "massmutual.com",
    "nationwide.com", "thrivent.com", "principal.com", "equitable.com", "nyl.com", "newyorklife.com",
    "prudential.com", "guardianlife.com", "kestra.com", "avantax.com", "hightoweradvisors.com", "captrust.com",
    "mercer.com", "creativeplanning.com", "marinerwealthadvisors.com",
}
SIGNATRY_DOMAINS = {"thesignatry.com", "signatry.com"}
SIGNATRY_INSIDER_LABELS = {"board member", "employer", "executive leader", "owner", "staff", "officer", "director",
                           "trustee", "employee"}
NO_OWNER = "None Assigned"  # shown wherever a contact has no HubSpot owner: pull, workbook, PDF badge, filename


def owner_display(props, inp=None):
    """The Contact Owner for display and filenames: HubSpot's owner_name, else the input file's owner, else NO_OWNER."""
    props = props or {}
    inp = inp or {}
    return (props.get("owner_name") or inp.get("owner") or "").strip() or NO_OWNER

# Model policy (September 2026 five-contact comparison): Fable 5.x default, Opus 5 acceptable / escalation,
# Haiku 4.5 only for thin or placeholder tiers, Sonnet 5 not for household work. Advisory – it flags, never blocks.
MODEL_POLICY_STANDARD = ("claude-fable-5", "claude-opus-5")
MODEL_POLICY_THIN = MODEL_POLICY_STANDARD + ("claude-haiku-4-5",)


def model_policy_flag(model_id, tier):
    """Return a data-quality flag string when model_id is off-policy for the tier, else None."""
    mid = (model_id or "").strip().lower()
    allowed = MODEL_POLICY_THIN if tier in ("thin", "placeholder") else MODEL_POLICY_STANDARD
    if any(mid.startswith(p) for p in allowed):
        return None
    return (f"Researched by an off-policy model ({model_id or 'unknown'}) – the skill's model policy is Fable 5.x by default, "
            f"Opus 5 for escalation{', Haiku 4.5 for thin/placeholder tiers' if tier in ('thin', 'placeholder') else ''}; "
            "review spouse and household conclusions with extra care (see SKILL.md, Model)")


ROLE_LABELS = {"employee", "board member", "officer", "founder", "owner", "staff", "volunteer leader",
               "director", "trustee", "executive"}
REFERRAL_LABELS = {"referred by", "referral", "referrer"}
PLACEHOLDER_FIRSTNAMES = {"husband", "wife", "spouse", "guest", "unknown", "tbd", "partner", "plus one", ""}


def now():
    return datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def skill_version():
    """Read this skill's own version from its SKILL.md frontmatter, so renderers never hand-carry a duplicate string."""
    path = os.path.join(os.path.dirname(__file__), "..", "SKILL.md")
    try:
        text = open(path, encoding="utf-8").read()
    except OSError:
        return None
    parts = text.split("---", 2)
    if len(parts) < 3:
        return None
    m = re.search(r"^version:\s*(\S+)", parts[1], re.M)
    return m.group(1) if m else None


def contact_path(run_dir, hid):
    return os.path.join(run_dir, "contacts", f"{hid}.json")


def load_json(path, default=None):
    if not os.path.exists(path):
        return default
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def save_json(path, obj):
    """Atomic write. The temp name carries the PID because parallel subagents all recount the shared manifest;
    a shared `manifest.json.tmp` let one writer's os.replace delete another's temp file mid-write (FileNotFoundError)."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = f"{path}.{os.getpid()}.tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(obj, f, indent=2, ensure_ascii=False)
    os.replace(tmp, path)


def load_manifest(run_dir):
    m = load_json(os.path.join(run_dir, "manifest.json"))
    if m is None:
        raise SystemExit(f"No manifest.json in {run_dir}; run init_run.py first.")
    return m


def save_manifest(run_dir, m):
    save_json(os.path.join(run_dir, "manifest.json"), m)


def list_contacts(run_dir):
    d = os.path.join(run_dir, "contacts")
    if not os.path.isdir(d):
        return []
    out = []
    for fn in sorted(os.listdir(d)):
        if fn.endswith(".json"):
            out.append(load_json(os.path.join(d, fn)))
    return out


def valid_id(hid):
    return isinstance(hid, str) and hid.isdigit()


def normalize_id(raw):
    """Return a clean string ID or None. Rejects scientific notation / decimals (Excel damage)."""
    if raw is None:
        return None
    s = str(raw).strip()
    if re.fullmatch(r"\d+", s):
        return s
    if re.fullmatch(r"\d+\.0+", s):
        return s.split(".")[0]
    return None  # e.g. "2.44057E+11"


def recount(run_dir):
    m = load_manifest(run_dir)
    cs = list_contacts(run_dir)
    c = {"total": len(cs), "init": 0, "hubspot": 0, "researched": 0, "rendered": 0, "rejected": m["counts"].get("rejected", 0)}
    for ct in cs:
        st = ct.get("status", {}).get("stages", {})
        c["init"] += 1
        if st.get("hubspot"):
            c["hubspot"] += 1
        if st.get("research"):
            c["researched"] += 1
        if st.get("rendered"):
            c["rendered"] += 1
    m["counts"] = c
    save_manifest(run_dir, m)
    return c


# ------------------------------------------------------------------ redaction
def screen_text(text, lexicon=True, insider=False):
    """Return (clean_text, hit_reason or None).

    The pattern screen (SSN, card, account, PIN, wire/ACH details, credential) always runs, with URLs removed first
    so a 14-digit press-release ID or an 18-digit EDGAR accession number in a source link is not mistaken for a card
    number. The health/hardship lexicon is meant for engagement bodies, where a note may record a donor's personal
    circumstances; pass lexicon=False for research prose, where 'medical' is a company's sector and 'Terminal' is a
    street name, not Restricted data about a person. Lexicon entries are stems matched at a word start, so 'diagnos'
    catches 'diagnosed' but 'rehab' does not fire inside 'prehabilitation'. insider=True adds the Board Confidential
    lexicon, for contacts who are Signatry staff or board members (see signatry_insider).
    """
    if not text:
        return text, None
    stripped = URL_RE.sub(" ", text)
    for pat, reason in RESTRICTED_PATTERNS:
        if pat.search(stripped):
            return "[redacted – possible Restricted content]", reason
    if lexicon:
        low = stripped.lower()
        for w in HEALTH_HARDSHIP_LEXICON:
            if re.search(r"\b" + re.escape(w), low):
                return "[redacted – possible Restricted content]", f"lexicon:{w}"
        if insider:
            for w in BOARD_CONFIDENTIAL_LEXICON:
                if re.search(r"\b" + re.escape(w) + r"\b", low):
                    return "[redacted – possible Board Confidential content]", f"board:{w}"
    return text, None


def signatry_insider(props, assoc):
    """Return 'staff', 'board', or None: is this contact a Signatry employee or board member?

    Staff: a Signatry email domain, or a company association named The Signatry carrying an employment-type label.
    Board: a Signatry-named company association labeled Board Member, or any association label containing 'board'
    on a Signatry-named company. Used to switch on the Board Confidential engagement screen (IT15).
    """
    props = props or {}
    assoc = assoc or {}
    email = (props.get("email") or "").strip().lower()
    domain = email.split("@")[-1] if "@" in email else ""
    for co in assoc.get("companies") or []:
        if "signatry" not in (co.get("name") or "").lower():
            continue
        labels = {str(l).lower() for l in (co.get("labels") or [])}
        if any("board" in l for l in labels) or "trustee" in labels:
            return "board"
        if labels & SIGNATRY_INSIDER_LABELS:
            return "staff"
    if domain in SIGNATRY_DOMAINS:
        return "staff"
    return None


def fix_name_case(name):
    """Title-case an owner name only when HubSpot returned it all-lowercase or all-uppercase ('dale armstrong'),
    leaving mixed-case names such as 'Jill McBroom' untouched."""
    if not name or "@" in name:
        return name
    s = str(name).strip()
    if s == s.lower() or s == s.upper():
        return " ".join(w[:1].upper() + w[1:].lower() for w in s.split())
    return s


def contains_characterization(text):
    if not text:
        return None
    low = text.lower()
    for w in CHARACTERIZATION_BLOCKLIST:
        if w in low:
            return w
    return None


# ------------------------------------------------------------------ naming
def safe(s):
    s = re.sub(r"[^A-Za-z0-9]+", "_", str(s or "")).strip("_")
    return s or "Unknown"


def profile_filename(ct, convention):
    hid = ct["hs_object_id"]
    props = (ct.get("hubspot") or {}).get("properties") or {}
    inp = ct.get("input") or {}
    first = safe(props.get("firstname") or inp.get("firstname"))
    last = safe(props.get("lastname") or inp.get("lastname"))
    owner = safe(owner_display(props, inp))
    if convention == 2:
        return f"{last}_{first}_{hid}_profile.pdf"
    if convention == 3:
        return f"{owner}_{last}_{first}_{hid}_profile.pdf"
    return f"{hid}_{first}_{last}_profile.pdf"


def read_table(path):
    """Read CSV or XLSX into list of dicts (header row = keys)."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
        return [dict(zip(hdr, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]
    with open(path, newline="", encoding="utf-8-sig") as f:
        return [r for r in csv.DictReader(f) if any((v or "").strip() for v in r.values())]
